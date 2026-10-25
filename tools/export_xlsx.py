#!/usr/bin/env python3
"""把 data.js 的資料表匯出成 .xlsx，上傳到 Google Drive 後轉成 Google 試算表就是網站的資料庫。

需要：Node.js（讀 sheet.js 的欄位定義）與 openpyxl（pip install openpyxl）
用法：python3 tools/export_xlsx.py [輸出路徑]，預設 dist/北越十日-資料庫.xlsx
"""
import json, pathlib, subprocess, sys
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = pathlib.Path(__file__).resolve().parent.parent
OUT = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / 'dist' / '北越十日-資料庫.xlsx'

js = """
global.self = global;
require(process.argv[1] + '/data.js');
const S = require(process.argv[1] + '/sheet.js');
process.stdout.write(JSON.stringify({ TABLES: S.TABLES, ENUMS: S.ENUMS, data: self.TRIP_DATA }));
"""
meta = json.loads(subprocess.run(['node', '-e', js, str(ROOT)], capture_output=True, check=True, text=True).stdout)
TABLES, ENUMS, DATA = meta['TABLES'], meta['ENUMS'], meta['data']['tables']

FONT = 'Arial'
BODY = Font(name=FONT, size=10)
HEAD = Font(name=FONT, size=10, bold=True, color='1F2A44')
HEAD_FILL = PatternFill('solid', fgColor='E6E9F2')
INPUT_FILL = PatternFill('solid', fgColor='FFF2CC')
TITLE = Font(name=FONT, size=14, bold=True)
THIN = Border(bottom=Side(style='thin', color='B7BDCC'))
WRAP = Alignment(wrap_text=True, vertical='top')
TOP = Alignment(vertical='top')

LONG = {'lead': 48, 'desc': 52, 'tip': 44, 'summary': 52, 'tips': 44, 'note': 36, 'detail': 36, 'next_action': 44, 'plan_b': 44,
        'text': 56, 'alt': 40, 'sources': 40, 'dish': 36, 'address': 36, 'maps_query': 32, 'hours': 30, 'price': 30, 'title': 30,
        'name_zh': 22, 'name_vn': 26, 'name': 26, 'value': 40, 'area': 30, 'check_in_out': 22, 'unverified_note': 28, 'credit': 22,
        'license_url': 28, 'source_url': 30, 'item': 24, 'group': 12, 'section': 14, 'label': 12, 'vn': 30, 'zh': 20, 'number': 18,
        'route_stops': 30, 'place_ids': 24, 'food_ids': 26, 'used_on': 14, 'best_day': 16, 'michelin': 14, 'when_label': 12}


def cell_value(v, typ):
    if typ == 'bool':
        return bool(v) if v not in ('', None) else False
    if typ == 'num':
        return v if isinstance(v, (int, float)) and not isinstance(v, bool) else (None if v in ('', None) else v)
    if typ and typ.startswith('enum:'):
        return ENUMS[typ[5:]].get(v, v) if v else None
    if isinstance(v, list):
        return ('\n' if typ == 'lines' else ', ').join(v)
    return None if v in ('', None) else str(v)


wb = Workbook()
guide = wb.active
guide.title = '說明'

# ---------- 預算總覽（公式即時計算，網站不讀這一頁） ----------
summary = wb.create_sheet('預算總覽')

sheet_of = {}
for t in TABLES:
    name = f"{t['label']} {t['key']}"
    ws = wb.create_sheet(name)
    sheet_of[t['key']] = name
    cols = t['cols']
    for ci, c in enumerate(cols, start=1):
        h = ws.cell(row=1, column=ci, value=f'{c[1]} {c[0]}')
        h.font, h.fill, h.border, h.alignment = HEAD, HEAD_FILL, THIN, Alignment(vertical='center', wrap_text=True)
        width = LONG.get(c[0], 12 if c[0].endswith('_id') or c[0] in ('date', 'due', 'time', 'end') else 11)
        ws.column_dimensions[get_column_letter(ci)].width = width
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = 'B2'
    rows = DATA.get(t['key'], [])
    for ri, r in enumerate(rows, start=2):
        for ci, c in enumerate(cols, start=1):
            typ = c[2] if len(c) > 2 else 'text'
            v = cell_value(r.get(c[0]), typ)
            if t['key'] == 'trip' and c[0] == 'value' and r.get('key') == 'rate_vnd_per_twd':
                v = float(v) if v else None
            cell = ws.cell(row=ri, column=ci, value=v)
            cell.font = BODY
            cell.alignment = WRAP if LONG.get(c[0], 0) >= 30 else TOP
            if typ == 'num':
                cell.number_format = '0.000000' if c[0] in ('lat', 'lng') else '#,##0'
            elif typ != 'bool' and not (t['key'] == 'trip' and c[0] == 'value' and r.get('key') == 'rate_vnd_per_twd'):
                cell.number_format = '@'
    # 下拉選單：列舉值與是／否；範圍拉到 500 列，之後新增的列也適用
    last = 500
    for ci, c in enumerate(cols, start=1):
        typ = c[2] if len(c) > 2 else 'text'
        col = get_column_letter(ci)
        if typ.startswith('enum:'):
            dv = DataValidation(type='list', formula1='"' + ','.join(ENUMS[typ[5:]].values()) + '"', allow_blank=True,
                                showErrorMessage=True, errorTitle='選項不在清單裡', error='請從下拉選單選一個值。')
            ws.add_data_validation(dv); dv.add(f'{col}2:{col}{last}')
        elif typ == 'bool':
            dv = DataValidation(type='list', formula1='"TRUE,FALSE"', allow_blank=True)
            ws.add_data_validation(dv); dv.add(f'{col}2:{col}{last}')
        elif typ in ('text', 'list', 'lines') or len(c) < 3:
            for rr in range(len(rows) + 2, last + 1):
                ws.cell(row=rr, column=ci).number_format = '@'
    note = t.get('note')
    if note:
        ws.cell(row=1, column=1).comment = Comment(note, '北越十日')
    if t['key'] in ('places', 'food', 'tips'):
        ws.cell(row=1, column=1).comment = Comment((note + '\n' if note else '') + '資料來源：2026/10/25–26 網路查證，每列的來源網址在「資料來源 sources」欄。', '北越十日')

# ---------- 說明頁 ----------
lines = [
    ('北越十日 2026 資料庫', TITLE),
    ('這份試算表是「北越十日」網站的資料來源。網站透過 Apps Script 讀取這裡的內容，改完存檔，手機下次連網開啟網站就會更新。', BODY),
    ('', BODY),
    ('怎麼編輯', HEAD),
    ('1. 每個分頁是一張資料表，第一列是欄位名稱，格式是「中文 代號」。中文可以改，最後的英文代號不要改，網站靠代號找欄位。', BODY),
    ('2. 分頁名稱也一樣是「中文 代號」，例如「每日 days」。不要改英文代號，也不要刪掉分頁。', BODY),
    ('3. 每一列的第一欄是代號（例如 day_id、place_id），不能重複。其他分頁用代號互相指到，例如「時間軸」的 place_ids 填「hoan-kiem, ngoc-son」。', BODY),
    ('4. 有下拉選單的欄位請從選單選；是／否欄位填 TRUE 或 FALSE。', BODY),
    ('5. 想新增一個景點：在「地點 places」加一列、取一個新的 place_id，再把它填到「時間軸 items」的 place_ids。', BODY),
    ('6. 改完可以從選單「北越十日 → 檢查資料」確認代號都找得到（需先完成 Apps Script 設定）。', BODY),
    ('', BODY),
    ('網站會寫回的分頁', HEAD),
    ('「行前清單 checklist」的 done 與 done_at：在網站勾選時自動寫入。', BODY),
    ('「記帳 expenses」：在網站按「記一筆」會新增一列；在這裡直接新增也會顯示在網站上。', BODY),
    ('「預算總覽」只在試算表裡用公式加總，網站不讀這一頁。', BODY),
    ('', BODY),
    ('隱私', HEAD),
    ('網站是公開的，這裡的內容會顯示在網站上。不要填機票票號、訂位代碼、護照號碼或付款資訊。', BODY),
    ('', BODY),
    ('連接網站（只需要做一次）', HEAD),
    ('1. 選單「擴充功能 → Apps Script」，把 GitHub repo 裡 apps-script/Code.gs 的內容整個貼上，存檔。', BODY),
    ('2. 左側「專案設定 → 指令碼屬性」新增 WRITE_KEY，值是你自己定的寫入密碼。', BODY),
    ('3. 「部署 → 新增部署作業 → 網頁應用程式」，執行身分選「我」，誰可以存取選「所有人」，按部署並授權。', BODY),
    ('4. 複製網頁應用程式網址（…/exec），貼到網站「出發前 → Google Sheet 同步」的資料 API 網址，寫入密碼填同一個 WRITE_KEY。', BODY),
]
for i, (text, font) in enumerate(lines, start=1):
    c = guide.cell(row=i, column=1, value=text)
    c.font = font
    c.alignment = Alignment(wrap_text=True, vertical='top')
guide.column_dimensions['A'].width = 110

# ---------- 預算總覽 ----------
BUD = f"'{sheet_of['budget']}'"
EXP = f"'{sheet_of['expenses']}'"
TRIP = f"'{sheet_of['trip']}'"
rate_row = 2 + [r['key'] for r in DATA['trip']].index('rate_vnd_per_twd')
summary['A1'] = '預算總覽'; summary['A1'].font = TITLE
summary['A2'] = '估算來自「預算 budget」分頁，實際花費來自「記帳 expenses」分頁，都是公式，新增列會自動算進來。黃色格子可以修改。'
summary['A2'].font = BODY; summary.merge_cells('A2:F2'); summary['A2'].alignment = Alignment(wrap_text=True)
summary.row_dimensions[2].height = 30
summary['A4'] = '人數'; summary['B4'] = 1
summary['B4'].fill = INPUT_FILL
summary['C4'] = '每人估算會乘上這個數字，整團的項目不乘。'
summary['A5'] = '匯率（1 TWD 換多少 VND）'; summary['B5'] = f'={TRIP}!B{rate_row}'
summary['C5'] = '來自「旅程設定 trip」的 rate_vnd_per_twd。'
summary['B4'].number_format = '0'; summary['B5'].number_format = '#,##0'
heads = ['類別', '估算低（VND）', '估算高（VND）', '估算低（NT$）', '估算高（NT$）', '實際花費（NT$）']
for ci, h in enumerate(heads, start=1):
    c = summary.cell(row=7, column=ci, value=h); c.font, c.fill, c.border = HEAD, HEAD_FILL, THIN
cats = list(ENUMS['category'].values())
for i, cat in enumerate(cats):
    r = 8 + i
    summary.cell(row=r, column=1, value=cat)
    summary.cell(row=r, column=2, value=f'=SUMIFS({BUD}!$G:$G,{BUD}!$C:$C,$A{r},{BUD}!$F:$F,"每人")*$B$4+SUMIFS({BUD}!$G:$G,{BUD}!$C:$C,$A{r},{BUD}!$F:$F,"整團")')
    summary.cell(row=r, column=3, value=f'=SUMIFS({BUD}!$H:$H,{BUD}!$C:$C,$A{r},{BUD}!$F:$F,"每人")*$B$4+SUMIFS({BUD}!$H:$H,{BUD}!$C:$C,$A{r},{BUD}!$F:$F,"整團")')
    summary.cell(row=r, column=4, value=f'=IF($B$5=0,0,B{r}/$B$5)')
    summary.cell(row=r, column=5, value=f'=IF($B$5=0,0,C{r}/$B$5)')
    summary.cell(row=r, column=6, value=f'=SUMIFS({EXP}!$F:$F,{EXP}!$D:$D,$A{r})+IF($B$5=0,0,SUMIFS({EXP}!$E:$E,{EXP}!$D:$D,$A{r},{EXP}!$F:$F,"")/$B$5)')
tot = 8 + len(cats)
summary.cell(row=tot, column=1, value='合計').font = HEAD
for ci in range(2, 7):
    col = get_column_letter(ci)
    c = summary.cell(row=tot, column=ci, value=f'=SUM({col}8:{col}{tot - 1})'); c.font = HEAD
for r in range(8, tot + 1):
    for ci in range(2, 7):
        summary.cell(row=r, column=ci).number_format = '#,##0;(#,##0);"-"'
for r in range(4, tot + 1):
    for ci in range(1, 7):
        cell = summary.cell(row=r, column=ci)
        if cell.font != HEAD and cell.font != TITLE: cell.font = Font(name=FONT, size=10, bold=cell.font.bold if cell.font else False)
summary.cell(row=tot + 2, column=1, value='不含機票與已付款的住宿；實際花費的台幣金額以記帳時的匯率換算。').font = BODY
summary.column_dimensions['A'].width = 26
for col in 'BCDEF': summary.column_dimensions[col].width = 16
summary.column_dimensions['C'].width = 18

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print('saved', OUT, OUT.stat().st_size, 'bytes')
